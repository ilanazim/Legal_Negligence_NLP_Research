import docx
from docx.document import Document
from docx.oxml.table import CT_Tbl
from docx.oxml.text.paragraph import CT_P
from docx.table import _Cell, Table
from docx.text.paragraph import Paragraph


def iter_block_items(parent):
    """
    Yield each paragraph and table child within *parent*, in document order.
    Each returned value is an instance of either Table or Paragraph. *parent*
    would most commonly be a reference to a main Document object, but
    also works for a _Cell object, which itself can contain paragraphs and tables.
    """
    if isinstance(parent, Document):
        parent_elm = parent.element.body
    elif isinstance(parent, _Cell):
        parent_elm = parent._tc
    else:
        raise ValueError("something's not right")

    for child in parent_elm.iterchildren():
        if isinstance(child, CT_P):
            yield Paragraph(child, parent)
        elif isinstance(child, CT_Tbl):
            yield Table(child, parent)

def convert_docx_to_txt(file_path, save_path = None, add_header_tag = True):    
    # Open the document
    # (Note: docx.Document is different from docx.document.Document)
    document = docx.Document(file_path)
    file_text = ''

    start_of_case = True
    doc_section = 0
    
    # Iterate over each item in the body
    for item in iter_block_items(document):

        # If at start of case, store the header (case name)
        # as the first entry. Must access 'document.sections[].header' to do so.
        if start_of_case:
            header = document.sections[doc_section].header
            case_title = header.tables[0].cell(1, 0).text
            file_text += case_title.strip() + '\n'
            doc_section += 2
            start_of_case = False
        
        # If item is a PARAGRAPH
        if isinstance(item, Paragraph):
            header = False
            for run in item.runs:
                if run.bold:
                    if run.text.strip() == item.text.strip():
                        if len(run.text.split()) < 12:
                            if item.text.strip() != 'End of Document':
                                header = True
            # Skip over any empty lines
            if len(item.text.strip()) > 0:
                if header and add_header_tag:
                    file_text += '<header>' + item.text.strip() + '</header>' + '\n'
                else:
                    file_text += item.text.strip() + '\n'

            # Check if we are at the last line of a case
            # to make sure we insert the header of the following case
            if item.text.strip() == 'End of Document':
                start_of_case = True 

            
              
        # If item is a TABLE
        elif isinstance(item, Table): 
            # We group everything across a row as one "sentence" in our .txt file
            # Currently have no better ideas. This seems to be the most common format of tables in our .docx files
            for row in item.rows:
                text = [cell.text.strip() for cell in row.cells]
                file_text += ' '.join(text) + '\n'
            
    # Save as .txt (if given a path)
    if save_path:
        with open(save_path, 'w') as out_file:
            out_file.write(file_text)
    else:
        # Else return the contents of the document
        return file_text